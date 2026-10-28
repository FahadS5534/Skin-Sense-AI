import os
import glob
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def markdown_to_docx(md_path: str, docx_path: str):
    doc = Document()

    # Set basic margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Filter separator row (e.g. |---|---|)
        cleaned_rows = []
        for r in table_rows:
            if all(set(c.strip()).issubset({"-", "|", ":", " "}) for c in r):
                continue
            cleaned_rows.append(r)

        if not cleaned_rows:
            in_table = False
            table_rows = []
            return

        cols_count = max(len(r) for r in cleaned_rows)
        tbl = doc.add_table(rows=len(cleaned_rows), cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.style = "Table Grid"

        for r_idx, row_cells in enumerate(cleaned_rows):
            for c_idx, cell_text in enumerate(row_cells):
                if c_idx < cols_count:
                    cell = tbl.cell(r_idx, c_idx)
                    cell.text = cell_text.strip()
                    if r_idx == 0:
                        # Header row styling
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.bold = True
                                run.font.size = Pt(9.5)
                    else:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(9)

        doc.add_paragraph()  # spacer
        in_table = False
        table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle Table
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            table_rows.append(cells)
            continue
        elif in_table:
            flush_table()

        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            h.runs[0].font.name = "Georgia"
            h.runs[0].font.color.rgb = RGBColor(43, 37, 35)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            h.runs[0].font.name = "Georgia"
            h.runs[0].font.color.rgb = RGBColor(217, 114, 107)
        elif stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
            h.runs[0].font.name = "Georgia"
            h.runs[0].font.color.rgb = RGBColor(59, 51, 48)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("---"):
            continue
        else:
            p = doc.add_paragraph(stripped)
            for run in p.runs:
                run.font.size = Pt(10.5)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Generated: {docx_path}")


def main():
    doc_dir = os.path.dirname(__file__)
    md_files = glob.glob(os.path.join(doc_dir, "**/*.md"), recursive=True)

    print(f"Found {len(md_files)} Markdown source documents to convert to Word (.docx)...")
    for md in md_files:
        docx = md[:-3] + ".docx"
        try:
            markdown_to_docx(md, docx)
        except Exception as e:
            print(f"Error converting {md}: {e}")

    print("All Word (.docx) documents generated successfully!")


if __name__ == "__main__":
    main()
